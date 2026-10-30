"""Convert reports/REPORT.md into a formatted Word document (reports/REPORT.docx)."""
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from config import REPORTS_DIR, ROOT

BRAND = RGBColor(0xE3, 0x1C, 0x5F)


def add_runs(p, text):
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if part.startswith("**"):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif part.startswith("*") and len(part) > 2:
            p.add_run(part[1:-1]).italic = True
        elif part:
            p.add_run(part)


def main():
    lines = (REPORTS_DIR / "REPORT.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    i, in_code, code = 0, False, []
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            if in_code:
                p = doc.add_paragraph(); r = p.add_run("\n".join(code)); r.font.name = "Consolas"; r.font.size = Pt(9)
                code = []
            in_code = not in_code
        elif in_code:
            code.append(ln)
        elif ln.startswith("# "):
            h = doc.add_heading(ln[2:], 0); h.runs[0].font.color.rgb = BRAND
        elif ln.startswith("## "):
            doc.add_heading(ln[3:], 1)
        elif ln.startswith("### "):
            doc.add_heading(ln[4:], 2)
        elif ln.startswith("!["):
            path = re.search(r"\((.+?)\)", ln).group(1).replace("../", "")
            doc.add_picture(str(ROOT / path), width=Inches(6.3))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-| ]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
            for r, row in enumerate(rows):
                for c, val in enumerate(row):
                    cell = t.cell(r, c); cell.text = ""
                    add_runs(cell.paragraphs[0], val)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9); run.bold = run.bold or r == 0
            doc.add_paragraph()
            continue
        elif ln.startswith("- "):
            add_runs(doc.add_paragraph(style="List Bullet"), ln[2:])
        elif ln.startswith("> "):
            p = doc.add_paragraph(); add_runs(p, ln[2:]); p.paragraph_format.left_indent = Inches(.3)
        elif ln.strip() and ln.strip() != "---":
            add_runs(doc.add_paragraph(), ln)
        i += 1
    out = REPORTS_DIR / "REPORT.docx"
    doc.save(out)
    print("Saved", out)


if __name__ == "__main__":
    main()
